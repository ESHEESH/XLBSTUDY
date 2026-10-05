"""Stateless AI backend: file/URL/text in, flashcards out. No DB access."""
import io
import os
import time
from collections import defaultdict, deque
from typing import Literal

import jwt
from dotenv import load_dotenv
from fastapi import Depends, FastAPI, File, Header, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse, StreamingResponse
from google import genai
from google.genai import types
from pydantic import BaseModel, Field, field_validator, model_validator

load_dotenv()
MODEL = os.getenv("GEMINI_MODEL", "gemini-2.5-flash")
SUPABASE_URL = os.getenv("SUPABASE_URL", "").rstrip("/")
MAX_BYTES = 20 * 1024 * 1024  # inline limit; ponytail: use Gemini Files API above this
client = genai.Client(api_key=os.getenv("GEMINI_API_KEY"))
jwks = jwt.PyJWKClient(f"{SUPABASE_URL}/auth/v1/.well-known/jwks.json") if SUPABASE_URL else None

app = FastAPI(docs_url=None, redoc_url=None, openapi_url=None)  # don't advertise the API surface
app.add_middleware(
    CORSMiddleware,
    allow_origins=os.getenv("ALLOWED_ORIGINS", "http://localhost:3000").split(","),
    allow_methods=["GET", "POST"],
    allow_headers=["Authorization", "Content-Type"],
)


@app.middleware("http")
async def cap_body(req, call_next):
    # ponytail: trusts Content-Length; chunked uploads need a proxy-level limit on the host
    if int(req.headers.get("content-length") or 0) > MAX_BYTES + 1_000_000:
        return JSONResponse({"detail": "Request too large"}, 413)
    return await call_next(req)


def current_user(authorization: str = Header("")) -> str:
    token = authorization.removeprefix("Bearer ").strip()
    if not jwks or not token:
        raise HTTPException(401, "Missing token")
    try:
        key = jwks.get_signing_key_from_jwt(token).key
        return jwt.decode(token, key, algorithms=["ES256", "RS256"], audience="authenticated")["sub"]
    except Exception:
        raise HTTPException(401, "Invalid token")


_hits: dict[tuple[str, str], deque] = defaultdict(deque)


def limit(bucket: str, n: int, per: int = 60):
    """Per-user sliding window. ponytail: in-process, so one worker; use Redis if you scale out."""

    def dep(uid: str = Depends(current_user)) -> str:
        q, now = _hits[(bucket, uid)], time.monotonic()
        while q and now - q[0] > per:
            q.popleft()
        if len(q) >= n:
            raise HTTPException(429, "Too many requests", headers={"Retry-After": str(per)})
        q.append(now)
        return uid

    return dep


gen_limit = limit("generate", 10)  # each call spends Gemini quota
tutor_limit = limit("tutor", 30)


class Card(BaseModel):
    kind: Literal["basic", "cloze"]
    front: str = Field(max_length=4000)
    back: str = Field("", max_length=4000)

    @model_validator(mode="after")
    def has_answer(self):
        if not self.front.strip():
            raise ValueError("card has no front")
        if self.kind == "basic" and not self.back.strip():
            raise ValueError("basic card has no answer")
        if self.kind == "cloze" and "{{c" not in self.front:
            raise ValueError("cloze card has no blank")
        return self


class Deck(BaseModel):
    title: str = Field(max_length=200)
    cards: list[Card] = Field(max_length=300)

    @field_validator("cards")
    @classmethod
    def non_empty(cls, v):
        if not v:
            raise ValueError("no cards")
        return v


PROMPT = (
    "Create study flashcards from the material. Return a short deck title and a mix of "
    "'basic' cards (front = question, back = answer) and 'cloze' cards (front uses Anki "
    "syntax like 'The {{c1::mitochondria}} makes ATP', back empty). One fact per card, "
    "self-contained, no trivia. Use the material's language."
)


def generate(parts: list) -> Deck:
    res = client.models.generate_content(
        model=MODEL,
        contents=[*parts, PROMPT],
        config=types.GenerateContentConfig(
            response_mime_type="application/json",
            # Gemini's schema can't express our validators; they run on parse below
            response_schema=Deck,
        ),
    )
    try:
        return Deck.model_validate_json(res.text)
    except Exception:
        raise HTTPException(502, "Model returned unusable cards, try again")


def extract_office(name: str, data: bytes) -> str:
    if name.endswith(".docx"):
        from docx import Document

        return "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)
    from pptx import Presentation

    return "\n".join(
        sh.text_frame.text
        for slide in Presentation(io.BytesIO(data)).slides
        for sh in slide.shapes
        if sh.has_text_frame
    )


NATIVE = {  # extension -> mime, sent straight to Gemini
    ".pdf": "application/pdf",
    ".mp3": "audio/mp3",
    ".m4a": "audio/mp4",
    ".webm": "audio/webm",
    ".wav": "audio/wav",
}


@app.get("/health")
def health():
    return "ok"


@app.post("/generate/file")
def generate_file(file: UploadFile = File(...), _: str = Depends(gen_limit)):
    data = file.file.read(MAX_BYTES + 1)
    if len(data) > MAX_BYTES:
        raise HTTPException(413, "File over 20 MB")
    name = (file.filename or "").lower()
    ext = os.path.splitext(name)[1]
    if ext in NATIVE:
        return generate([types.Part.from_bytes(data=data, mime_type=NATIVE[ext])])
    if ext in (".docx", ".pptx"):
        return generate([extract_office(name, data)])
    if ext == ".txt":
        return generate([data.decode("utf-8", "replace")])
    raise HTTPException(415, "Unsupported file type")


class UrlIn(BaseModel):
    url: str


@app.post("/generate/youtube")
def generate_youtube(body: UrlIn, _: str = Depends(gen_limit)):
    if not body.url.startswith(("https://www.youtube.com/", "https://youtube.com/", "https://youtu.be/")):
        raise HTTPException(400, "Not a YouTube URL")
    return generate([types.Part(file_data=types.FileData(file_uri=body.url))])


class TextIn(BaseModel):
    text: str


@app.post("/generate/text")
def generate_text(body: TextIn, _: str = Depends(gen_limit)):
    if len(body.text) > 200_000:
        raise HTTPException(413, "Text too long")
    return generate([body.text])


class Msg(BaseModel):
    role: Literal["user", "model"]
    text: str = Field(max_length=4000)


class TutorIn(BaseModel):
    card: Card
    messages: list[Msg] = Field([], max_length=30)
    mode: Literal["eli5", "steps", "free"] = "free"


MODES = {
    "eli5": "Explain this flashcard's concept like I'm 5. Be brief.",
    "steps": "Explain this flashcard's concept step by step. Be brief.",
    "free": "You are a concise study tutor. Answer the student's questions.",
}


@app.post("/tutor")
def tutor(body: TutorIn, _: str = Depends(tutor_limit)):
    contents = [
        types.Content(role=m.role, parts=[types.Part(text=m.text)])
        for m in body.messages
    ] or [types.Content(role="user", parts=[types.Part(text=MODES[body.mode])])]

    def stream():
        for chunk in client.models.generate_content_stream(
            model=MODEL,
            contents=contents,
            config=types.GenerateContentConfig(
                system_instruction=f"{MODES[body.mode]}\nCard front: {body.card.front}\nCard back: {body.card.back}"
            ),
        ):
            if chunk.text:
                yield f"data: {chunk.text.replace(chr(10), chr(92) + 'n')}\n\n"

    return StreamingResponse(stream(), media_type="text/event-stream")


if __name__ == "__main__":
    # Phase 3 check: a card with no answer is rejected.
    for bad in ({"kind": "basic", "front": "Q"}, {"kind": "cloze", "front": "no blank"}):
        try:
            Card(**bad)
            raise SystemExit(f"accepted bad card {bad}")
        except ValueError:
            pass
    Card(kind="cloze", front="The {{c1::heart}} pumps blood")
    # tutor input is bounded and role-checked
    for bad in ({"card": {"kind": "basic", "front": "Q", "back": "A"}, "messages": [{"role": "system", "text": "x"}]},
                {"card": {"kind": "basic", "front": "Q", "back": "A"}, "messages": [{"role": "user", "text": "x" * 4001}]}):
        try:
            TutorIn(**bad)
            raise SystemExit("accepted bad tutor input")
        except ValueError:
            pass
    print("ok")
